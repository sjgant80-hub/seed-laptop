# measure/readers.py — the mainstream readers of claim L3, each opening every file the office wrote.
#
#   python3 measure/readers.py <folder of written files> <ci|local>   → one JSON object on stdout:
#   { reader: { file: { "opened": true, "content": … } | { "opened": false, "error": "…" } } }
#
# ci    (GitHub's ubuntu runner): LibreOffice headless, poppler pdftotext, PyMuPDF, pypdf, python-docx, openpyxl,
#       Python's email package, icalendar, Python's csv
# local (this laptop):            PyMuPDF, pypdf, Python's email package, Python's csv
# A reader that is not installed is reported as missing, never skipped silently.
import json, os, subprocess, sys, csv, io, shutil, tempfile
import email, email.policy

folder, where = sys.argv[1], sys.argv[2]
files = sorted(os.listdir(folder))
out = {}

def each(reader, ext, fn):
    res = {}
    for f in files:
        if not f.endswith('.' + ext): continue
        try: res[f] = {'opened': True, 'content': fn(os.path.join(folder, f))}
        except Exception as e: res[f] = {'opened': False, 'error': (type(e).__name__ + ': ' + str(e))[:300]}
    out.setdefault(reader, {}).update(res)

def missing(reader, why): out[reader] = {'missing': why}

def pdf_pymupdf(p):
    import pymupdf
    d = pymupdf.open(p)
    return '\n'.join(pg.get_text() for pg in d)
def pdf_pypdf(p):
    import pypdf
    return '\n'.join(pg.extract_text() for pg in pypdf.PdfReader(p).pages)
def pdf_pdftotext(p):
    return subprocess.run(['pdftotext', '-enc', 'UTF-8', p, '-'], capture_output=True, text=True, check=True, timeout=60).stdout
def docx_python(p):
    import docx
    d = docx.Document(p)
    return {'paragraphs': [x.text for x in d.paragraphs], 'cells': [c.text for t in d.tables for r in t.rows for c in r.cells]}
def xlsx_openpyxl(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True)
    return {ws.title: [[c for c in row] for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
def eml_python(p):
    m = email.message_from_bytes(open(p, 'rb').read(), policy=email.policy.default)
    return {'subject': str(m['subject']), 'from': str(m['from']), 'to': str(m['to']), 'date': str(m['date']), 'text': m.get_content()}
def ics_icalendar(p):
    import icalendar
    cal = icalendar.Calendar.from_ical(open(p, 'rb').read())
    return [{'uid': str(e.get('uid')), 'summary': str(e.get('summary')), 'location': str(e.get('location') or ''), 'description': str(e.get('description') or ''),
             'start': e.get('dtstart').dt.strftime('%Y%m%dT%H%M%SZ'), 'end': e.get('dtend').dt.strftime('%Y%m%dT%H%M%SZ')} for e in cal.walk('VEVENT')]
def csv_python(p):
    return list(csv.reader(io.StringIO(open(p, encoding='utf-8', newline='').read())))
def lo(p, target):
    d = tempfile.mkdtemp()
    r = subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', target, '--outdir', d, os.path.abspath(p)], capture_output=True, text=True, timeout=180)
    produced = [x for x in os.listdir(d)]
    if not produced: raise RuntimeError('LibreOffice did not open it: ' + (r.stderr or r.stdout)[-200:])
    return open(os.path.join(d, produced[0]), encoding='utf-8-sig', errors='replace').read()

for name, ext, fn in [('PyMuPDF', 'pdf', pdf_pymupdf), ('pypdf', 'pdf', pdf_pypdf)]:
    try: __import__('pymupdf' if name == 'PyMuPDF' else 'pypdf'); each(name, ext, fn)
    except ImportError: missing(name, 'not installed')
each('Python email', 'eml', eml_python)
each('Python csv', 'csv', csv_python)
if where == 'ci':
    if shutil.which('pdftotext'): each('poppler pdftotext', 'pdf', pdf_pdftotext)
    else: missing('poppler pdftotext', 'not installed')
    for name, ext, fn, mod in [('python-docx', 'docx', docx_python, 'docx'), ('openpyxl', 'xlsx', xlsx_openpyxl, 'openpyxl'), ('icalendar', 'ics', ics_icalendar, 'icalendar')]:
        try: __import__(mod); each(name, ext, fn)
        except ImportError: missing(name, 'not installed')
    if shutil.which('soffice'):
        each('LibreOffice (DOCX)', 'docx', lambda p: lo(p, 'txt:Text (encoded):UTF8'))
        each('LibreOffice (XLSX)', 'xlsx', lambda p: lo(p, 'csv:Text - txt - csv (StarCalc):44,34,76'))
        each('LibreOffice (CSV)', 'csv', lambda p: lo(p, 'csv:Text - txt - csv (StarCalc):44,34,76'))
    else: missing('LibreOffice', 'not installed')
print(json.dumps(out, ensure_ascii=True, default=str))
