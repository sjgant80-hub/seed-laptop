# measure/corpus.py — the L2a corpus: real files written by mainstream writers, from measure/corpus-content.json.
# Run on GitHub's ubuntu runner by .github/workflows/corpus.yml (never on the owner's laptop); its output is committed
# once, pinned by SHA-256 in seal 2, and never regenerated for a measurement.
#
#   python3 measure/corpus.py <outdir>
#
# Writers: python-docx, openpyxl, ReportLab, PyMuPDF, LibreOffice (headless), Google Chrome (headless print to PDF),
# Python's email package, icalendar, vobject, Pillow, zipfile, Python's csv. Each file's truth (what its writer was
# given) goes to <outdir>/truth.json; the writers' names and versions to <outdir>/writers.json.
import csv, io, json, os, subprocess, sys, zipfile, shutil, hashlib
from email.message import EmailMessage
from email import policy

out = os.path.abspath(sys.argv[1])
os.makedirs(out, exist_ok=True)
C = json.load(open(os.path.join(os.path.dirname(__file__), 'corpus-content.json'), encoding='utf-8'))
truth, writers = {}, {}
tmp = os.path.join(out, '_tmp'); os.makedirs(tmp, exist_ok=True)

def soffice(args, cwd):
    r = subprocess.run(['soffice', '--headless', '--norestore'] + args, cwd=cwd, timeout=180, capture_output=True, text=True)
    print('soffice', ' '.join(args), '->', r.returncode, r.stdout.strip()[-200:], r.stderr.strip()[-200:])

def words(texts):
    return ' '.join(texts)

letter, sheet, notes = C['letter'], C['sheet'], C['notes']
table_text = [c for row in letter['table'] for c in row]

# ── DOCX ──
import docx as pydocx
d = pydocx.Document()
for p in letter['paragraphs']: d.add_paragraph(p)
t = d.add_table(rows=len(letter['table']), cols=len(letter['table'][0]))
for r, row in enumerate(letter['table']):
    for c, v in enumerate(row): t.cell(r, c).text = v
d.core_properties.author = 'Example Author'
d.save(os.path.join(out, 'letter-pythondocx.docx'))
truth['letter-pythondocx.docx'] = {'type': 'docx', 'paragraphs': letter['paragraphs'], 'cells': table_text}
writers['letter-pythondocx.docx'] = 'python-docx ' + pydocx.__version__ if hasattr(pydocx, '__version__') else 'python-docx'

html = '<!doctype html><html><head><meta charset="utf-8"><title>Letter</title></head><body>' + ''.join('<p>' + p + '</p>' for p in letter['paragraphs']) + '<table border="1">' + ''.join('<tr>' + ''.join('<td>' + c + '</td>' for c in row) + '</tr>' for row in letter['table']) + '</table></body></html>'
open(os.path.join(tmp, 'letter.html'), 'w', encoding='utf-8').write(html)
soffice(['--infilter=HTML (StarWriter)', '--convert-to', 'docx:MS Word 2007 XML', '--outdir', tmp, 'letter.html'], tmp)
shutil.copy(os.path.join(tmp, 'letter.docx'), os.path.join(out, 'letter-libreoffice.docx'))
truth['letter-libreoffice.docx'] = {'type': 'docx', 'paragraphs': letter['paragraphs'], 'cells': table_text}
writers['letter-libreoffice.docx'] = 'LibreOffice Writer (headless, from HTML)'

# ── XLSX ──
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = sheet['name']
for row in sheet['rows']: ws.append(row)
wb.properties.creator = 'Example Author'
wb.save(os.path.join(out, 'orders-openpyxl.xlsx'))
cells = {}
for r, row in enumerate(sheet['rows']):
    for c, v in enumerate(row): cells[openpyxl.utils.get_column_letter(c + 1) + str(r + 1)] = v
truth['orders-openpyxl.xlsx'] = {'type': 'xlsx', 'sheet': sheet['name'], 'cells': cells}
writers['orders-openpyxl.xlsx'] = 'openpyxl ' + openpyxl.__version__

os.makedirs(os.path.join(tmp, 'lo'), exist_ok=True)
shutil.copy(os.path.join(out, 'orders-openpyxl.xlsx'), os.path.join(tmp, 'orders.xlsx'))
soffice(['--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', os.path.join(tmp, 'lo'), 'orders.xlsx'], tmp)
shutil.copy(os.path.join(tmp, 'lo', 'orders.xlsx'), os.path.join(out, 'orders-libreoffice.xlsx'))
truth['orders-libreoffice.xlsx'] = {'type': 'xlsx', 'sheet': sheet['name'], 'cells': cells}
writers['orders-libreoffice.xlsx'] = 'LibreOffice Calc (headless, re-saved)'

# ── PDF ──
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
import reportlab
cv = canvas.Canvas(os.path.join(out, 'letter-reportlab.pdf'), pagesize=A4)
y = 800
for p in letter['paragraphs']:
    line = ''
    for w in p.split(' '):
        if cv.stringWidth(line + ' ' + w if line else w, 'Helvetica', 11) > 480: cv.drawString(56, y, line); y -= 15; line = w
        else: line = line + ' ' + w if line else w
    cv.drawString(56, y, line); y -= 22
cv.save()
truth['letter-reportlab.pdf'] = {'type': 'pdf', 'text': words(letter['paragraphs'])}
writers['letter-reportlab.pdf'] = 'ReportLab ' + reportlab.Version

import pymupdf
doc = pymupdf.open(); page = doc.new_page(); y = 72
for p in notes:
    rect = pymupdf.Rect(56, y, 540, y + 60)
    page.insert_textbox(rect, p, fontsize=11, fontname='helv'); y += 64
doc.save(os.path.join(out, 'notes-pymupdf.pdf'))
truth['notes-pymupdf.pdf'] = {'type': 'pdf', 'text': words(notes)}
writers['notes-pymupdf.pdf'] = 'PyMuPDF ' + pymupdf.VersionBind

shutil.copy(os.path.join(out, 'letter-pythondocx.docx'), os.path.join(tmp, 'letterd.docx'))
soffice(['--convert-to', 'pdf', '--outdir', tmp, 'letterd.docx'], tmp)
shutil.copy(os.path.join(tmp, 'letterd.pdf'), os.path.join(out, 'letter-libreoffice.pdf'))
truth['letter-libreoffice.pdf'] = {'type': 'pdf', 'text': words(letter['paragraphs'] + table_text)}
writers['letter-libreoffice.pdf'] = 'LibreOffice Writer (headless, DOCX to PDF)'

open(os.path.join(tmp, 'notes.html'), 'w', encoding='utf-8').write('<!doctype html><meta charset="utf-8"><body style="font:14px Arial">' + ''.join('<p>' + p + '</p>' for p in notes))
chrome = shutil.which('google-chrome') or shutil.which('chromium') or shutil.which('chromium-browser')
subprocess.run([chrome, '--headless=new', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer', '--print-to-pdf=' + os.path.join(out, 'notes-chrome.pdf'), 'file://' + os.path.abspath(os.path.join(tmp, 'notes.html'))], check=True, timeout=120, capture_output=True)
truth['notes-chrome.pdf'] = {'type': 'pdf', 'text': words(notes)}
writers['notes-chrome.pdf'] = subprocess.run([chrome, '--version'], capture_output=True, text=True).stdout.strip()

# ── EML ──
for i, m in enumerate(C['mail']):
    msg = EmailMessage(policy=policy.SMTP)
    msg['From'] = m['from']; msg['To'] = m['to']; msg['Subject'] = m['subject']; msg['Date'] = m['date']
    msg['Message-ID'] = '<corpus-%d@example.org>' % (i + 1)
    msg.set_content(m['text'])
    att = []
    if m.get('html'): msg.add_alternative('<p>' + m['text'].replace('\n', '<br>') + '</p>', subtype='html')
    if m.get('attachment'):
        data = m['attachment']['text'].encode('utf-8')
        msg.add_attachment(data, maintype='text', subtype='plain', filename=m['attachment']['name'])
        att.append({'name': m['attachment']['name'], 'sha256': hashlib.sha256(data).hexdigest()})
    name = 'mail-%d.eml' % (i + 1)
    open(os.path.join(out, name), 'wb').write(msg.as_bytes())
    truth[name] = {'type': 'eml', 'subject': m['subject'], 'from': m['from'], 'to': m['to'], 'date': m['date'], 'text': m['text'], 'attachments': att}
    writers[name] = "Python email package (EmailMessage, policy.SMTP)"

# ── ICS ──
import icalendar
from datetime import datetime, timezone
cal = icalendar.Calendar(); cal.add('prodid', '-//Example//Corpus//EN'); cal.add('version', '2.0')
for e in C['events']:
    ev = icalendar.Event(); ev.add('uid', e['uid']); ev.add('summary', e['summary']); ev.add('location', e['location'])
    ev.add('dtstart', datetime.strptime(e['start'], '%Y%m%dT%H%M%SZ').replace(tzinfo=timezone.utc))
    ev.add('dtend', datetime.strptime(e['end'], '%Y%m%dT%H%M%SZ').replace(tzinfo=timezone.utc))
    ev.add('dtstamp', datetime(2026, 10, 29, 12, 0, 0, tzinfo=timezone.utc))
    cal.add_component(ev)
open(os.path.join(out, 'diary.ics'), 'wb').write(cal.to_ical())
truth['diary.ics'] = {'type': 'ics', 'events': C['events']}
writers['diary.ics'] = 'icalendar ' + icalendar.__version__

# ── VCF ──
import vobject
cards = []
for c in C['contacts']:
    v = vobject.vCard(); v.add('n').value = vobject.vcard.Name(family=c['family'], given=c['given']); v.add('fn').value = c['fn']
    e = v.add('email'); e.value = c['email']; e.type_param = 'INTERNET'
    t = v.add('tel'); t.value = c['tel']; t.type_param = 'CELL'
    if c['org']: v.add('org').value = [c['org']]
    if c['title']: v.add('title').value = c['title']
    cards.append(v.serialize())
open(os.path.join(out, 'contacts.vcf'), 'w', encoding='utf-8', newline='').write(''.join(cards))
truth['contacts.vcf'] = {'type': 'vcf', 'contacts': C['contacts']}
writers['contacts.vcf'] = 'vobject (vCard 3.0)'

# ── CSV ──
buf = io.StringIO(); csv.writer(buf).writerows(sheet['rows'])
open(os.path.join(out, 'orders-python.csv'), 'w', encoding='utf-8', newline='').write(buf.getvalue())
truth['orders-python.csv'] = {'type': 'csv', 'rows': [[str(v) for v in row] for row in sheet['rows']]}
writers['orders-python.csv'] = 'Python csv'
soffice(['--convert-to', 'csv:Text - txt - csv (StarCalc):44,34,76', '--outdir', os.path.join(tmp, 'lo'), 'orders.xlsx'], tmp)
shutil.copy(os.path.join(tmp, 'lo', 'orders.csv'), os.path.join(out, 'orders-libreoffice.csv'))
truth['orders-libreoffice.csv'] = {'type': 'csv', 'rows': [[('TRUE' if v is True else 'FALSE' if v is False else str(v)) for v in row] for row in sheet['rows']]}
writers['orders-libreoffice.csv'] = 'LibreOffice Calc (headless, XLSX to CSV)'

# ── images ──
from PIL import Image, ImageDraw
import PIL
for im in C['images']:
    img = Image.new('RGB', (im['width'], im['height']), (47, 107, 79))
    dr = ImageDraw.Draw(img); dr.rectangle([10, 10, im['width'] - 10, im['height'] - 10], outline=(255, 255, 255), width=4)
    for ext, fmt in ([('png', 'PNG'), ('gif', 'GIF')] if im['name'] == 'logo' else [('jpg', 'JPEG'), ('webp', 'WEBP')]):
        name = im['name'] + '.' + ext
        img.save(os.path.join(out, name), fmt)
        truth[name] = {'type': 'jpeg' if ext == 'jpg' else ext, 'width': im['width'], 'height': im['height']}
        writers[name] = 'Pillow ' + PIL.__version__

# ── ZIP ──
with zipfile.ZipFile(os.path.join(out, 'drawings.zip'), 'w', compression=zipfile.ZIP_DEFLATED) as z:
    for f in C['archive']: z.writestr(f['name'], f['text'])
truth['drawings.zip'] = {'type': 'zip', 'entries': [{'name': f['name'], 'size': len(f['text'].encode('utf-8'))} for f in C['archive']]}
writers['drawings.zip'] = 'Python zipfile (deflate)'

shutil.rmtree(tmp)
json.dump(truth, open(os.path.join(out, 'truth.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
json.dump(writers, open(os.path.join(out, 'writers.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
print('corpus:', len(truth), 'files')
